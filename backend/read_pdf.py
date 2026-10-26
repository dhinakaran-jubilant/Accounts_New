"""
Project: LoanDesk
Author: Dhinakaran Sekar
Email: dhinakaran.s@jubilantenterprises.in
Date: 2026-04-30 18:41
Description: PDF statement parser extracting tabular transaction records using pdfplumber with OCR fallback support.
"""

import pdfplumber
import pandas as pd
import re
import os
import logging

logger = logging.getLogger(__name__)

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not installed. Word document extraction will be disabled.")

try:
    import pytesseract

    TESS_PATH = os.environ.get(
        "TESSERACT_PATH",
        r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    )
    if os.path.exists(TESS_PATH):
        pytesseract.pytesseract.tesseract_cmd = TESS_PATH
        TESSERACT_AVAILABLE = True
    else:
        TESSERACT_AVAILABLE = False
        logger.warning(
            f"Tesseract binary not found at '{TESS_PATH}'. "
            "OCR for scanned PDFs will be disabled. "
            "Set the TESSERACT_PATH environment variable to override."
        )
except ImportError:
    TESSERACT_AVAILABLE = False
    logger.warning("pytesseract not installed. OCR for scanned PDFs will be disabled.")


def _ocr_page(page):
    """Run OCR on a pdfplumber page. Returns extracted text or empty string."""
    if not TESSERACT_AVAILABLE:
        return ""
    try:
        img = page.to_image(resolution=300).original
        return pytesseract.image_to_string(img)
    except Exception as e:
        logger.warning(f"OCR failed on page: {e}")
        return ""


def clean_customer_name(raw_name):
    """
    Cleans customer name extracted from PDF / OCR:
    - Strips photo/avatar placeholders and icon noise (e.g. 'vr, Nolmage', 'Nolmage', 'No Image', 'No Photo')
    - Corrects common OCR character misrecognitions like ']' instead of 'J' (e.g. 'AF]' -> 'AFJ')
    - Strips unwanted trailing commas, punctuation, and whitespace
    """
    if not raw_name:
        return ""
    s = str(raw_name).strip()
    # Strip photo thumbnail/icon OCR artifacts and NoImage text
    s = re.sub(r'(?i)\b(?:vr|yr|v~|vi|iv|pic|photo|img|icon|thumbnail)\b\s*,?\s*(?:No\s*Image|NoImage|Nolmage|No\s*Photo)?', '', s)
    s = re.sub(r'(?i)\b(?:No\s*Image|NoImage|Nolmage|No\s*Photo|NoPhoto|S/o|W/o|D/o|C/o)\b', '', s)
    s = re.sub(r'(?i)\b(?:vr|yr|v~|vi|iv)\b\s*$', '', s)

    # Fix OCR misreading of 'J' as ']'
    s = re.sub(r'(?<=[A-Z])\]', 'J', s)
    s = re.sub(r'\](?=[A-Z])', 'J', s)
    s = re.sub(r'\b\](?=[a-zA-Z])', 'J', s)

    s = re.sub(r'\s+', ' ', s)
    s = re.sub(r'[,.\s\-_/]+$', '', s).strip()
    return s


def parse_ocr_schedule(schedule_text, expected_emi=None):
    """
    Robust parser for OCR amortization schedules where table cells/dates wrap across lines.
    Reconstructs dates (e.g. '21-' \n '09-' \n '2026' -> '21-09-2026') and matches EMI and Interest.
    """
    schedule = []
    lines = [l.strip() for l in schedule_text.splitlines() if l.strip()]

    inst_blocks = []
    current_block = []

    for line in lines:
        if re.match(r'^\d{1,3}\s+(?:\d{1,2}[-/]|\d{1,2}\s+)', line):
            if current_block:
                inst_blocks.append(current_block)
            current_block = [line]
        elif current_block:
            if line.startswith('http') or 'page' in line.lower() or 'finance' in line.lower():
                continue
            current_block.append(line)

    if current_block:
        inst_blocks.append(current_block)

    for block in inst_blocks:
        first_line = block[0]
        # Match sNo and date parts on first line
        m = re.match(r'^(\d+)\s+(\d{1,2})[-/](\d{1,2})?[-/]?(\d{4})?\s*(.*)$', first_line)
        if not m:
            continue
        s_no = int(m.group(1))
        part_day = m.group(2)
        part_month = m.group(3)
        part_year = m.group(4)
        rest_of_line = m.group(5)

        for sub_line in block[1:]:
            # Check combined month-year '09-2026'
            my_m = re.match(r'^(\d{1,2})[-/](\d{4})$', sub_line)
            if my_m:
                if not part_month:
                    part_month = my_m.group(1)
                if not part_year:
                    part_year = my_m.group(2)
                continue
            # Check month '09-' or '09'
            if not part_month:
                sub_m = re.match(r'^(\d{1,2})[-/]?$', sub_line)
                if sub_m:
                    part_month = sub_m.group(1)
                    continue
            # Check year '2026'
            if not part_year:
                sub_y = re.match(r'^(\d{4})$', sub_line)
                if sub_y:
                    part_year = sub_y.group(1)
                    continue

        if part_day and part_month and part_year:
            due_date = f'{int(part_day):02d}-{int(part_month):02d}-{part_year}'
        else:
            due_date = ''

        nums = [float(x.replace(',', '')) for x in re.findall(r'[\d,]+\.?\d*', rest_of_line)]

        emi = None
        interest = 0.0

        if len(nums) >= 4:
            # Check Principal (nums[2]) + Interest (nums[3]) == EMI (nums[1])
            if abs((nums[2] + nums[3]) - nums[1]) < 2.0:
                emi = nums[1]
                interest = nums[3]
            elif abs((nums[1] + nums[2]) - nums[0]) < 2.0:
                emi = nums[0]
                interest = nums[2]
            elif expected_emi and abs(nums[1] - expected_emi) < 2.0:
                emi = nums[1]
                interest = nums[3]
            else:
                emi = nums[1]
                interest = nums[3]
        elif len(nums) >= 3:
            if abs((nums[1] + nums[2]) - nums[0]) < 2.0:
                emi = nums[0]
                interest = nums[2]
            else:
                emi = nums[0]
                interest = nums[1]
        elif len(nums) >= 1:
            emi = nums[0]

        if emi is not None:
            schedule.append({
                'date': due_date,
                'amount': emi,
                'interest_amount': interest,
                'cheque_no': '',
                'received_date': None,
                'payment_date': None
            })

    return schedule


def _extract_pdf_data_original(pdf_path):
    output = {
        "loan_number": None,
        "primary_account": None,
        "client_account": None,
        "address": None,
        "contact": None,
        "loan_date": None,
        "loan_amount": None,
        "total_repayment": None,
        "interest_rate": None,
        "emi": None,
        "total_installments": None,
        "table": []
    }

    with pdfplumber.open(pdf_path) as pdf:
        first_page = pdf.pages[0]
        text = first_page.extract_text() or ""

        # ------------------------------------------------------------------
        # Check if PDF is an image/scanned document and run OCR if available
        # ------------------------------------------------------------------
        is_scanned = not text.strip()
        if is_scanned:
            text = _ocr_page(first_page)
            if not text.strip():
                raise ValueError(
                    "The uploaded PDF appears to be a scanned image and no text could be "
                    "extracted. Please upload a text-searchable PDF, or ensure Tesseract "
                    "OCR is installed and the TESSERACT_PATH environment variable is set."
                )

        # 1. Loan Number
        loan_match = re.search(r"Loan Number\s+([$A-Z0-9]+)", text)
        if loan_match:
            output["loan_number"] = loan_match.group(1).replace('$', 'S')

        # 2. Branch (Primary Account Name)
        branch_match = re.search(r"Branch\s+(.+?)\s+Customer", text, re.S)
        if branch_match:
            raw = branch_match.group(1).strip()
            raw = re.sub(r"\s+", " ", raw)
            output["primary_account"] = raw.upper()

        # 3. Customer Details (Name, Address & Contact)
        cust_match = re.search(r"Customer Details\s*\n(.*?)(?:Agreement Details|Amortization Schedule|$)", text, re.S)
        if cust_match:
            cust_text = cust_match.group(1).strip()
            raw_lines = [line.strip() for line in cust_text.splitlines() if line.strip()]
            if raw_lines:
                # First line is name
                output["client_account"] = clean_customer_name(raw_lines[0])

                addr_lines = []
                contact_found = None
                for line in raw_lines[1:]:
                    cm = re.search(r"(?:Contact(?:\s*No)?|Mobile(?:\s*No)?|Phone(?:\s*No)?|Cell(?:\s*No)?)\s*[:.-]?\s*([+0-9\s,\-/]+)", line, re.I)
                    if cm:
                        contact_found = re.sub(r"\s+", "", cm.group(1)).strip()
                    else:
                        addr_lines.append(line)

                if addr_lines:
                    addr_combined = ", ".join(addr_lines)
                    addr_combined = re.sub(r"(\s*,\s*)+", ", ", addr_combined).strip(", ")
                    output["address"] = addr_combined
                if contact_found:
                    output["contact"] = contact_found
        else:
            cust_block = re.search(r"Customer Details\s*\n\s*([^\n]+)", text, re.S)
            if cust_block:
                output["client_account"] = clean_customer_name(cust_block.group(1))

        # 4. Agreement Table Extraction
        if not is_scanned:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    df = pd.DataFrame(table)
                    header_index = None
                    for i, row in df.iterrows():
                        row_clean = [
                            re.sub(r'\s+', ' ', str(x)).strip().lower()
                            for x in row if pd.notna(x) and str(x).strip()
                        ]
                        row_text = " ".join(row_clean)
                        if "loan date" in row_text and "loan amount" in row_text:
                            header_index = i
                            break

                    if header_index is not None:
                        df.columns = df.iloc[header_index]
                        df = df[header_index + 1:].reset_index(drop=True)

                        def clean_col(c, idx):
                            base = re.sub(r'\s+', ' ', str(c)).strip().lower()
                            return base if base and base != 'none' else f"col_{idx}"

                        df.columns = [clean_col(col, idx) for idx, col in enumerate(df.columns)]

                        if not df.empty:
                            row = df.iloc[0]
                            loan_date_col = next((c for c in df.columns if "loan date" in c), None)
                            loan_amt_col = next((c for c in df.columns if "loan amount" in c), None)
                            total_amt_col = next((c for c in df.columns if "total amount" in c), None)

                            if loan_date_col and loan_amt_col:
                                try:
                                    val_date = row.get(loan_date_col)
                                    val_date = val_date.iloc[0] if isinstance(val_date, pd.Series) else val_date
                                    output["loan_date"] = str(val_date).strip()

                                    val_amt = row.get(loan_amt_col)
                                    val_amt = val_amt.iloc[0] if isinstance(val_amt, pd.Series) else val_amt
                                    amt_str = str(val_amt).replace(",", "").strip()
                                    output["loan_amount"] = float(amt_str) if amt_str and amt_str != 'None' else None

                                    if total_amt_col:
                                        val_tot = row.get(total_amt_col)
                                        val_tot = val_tot.iloc[0] if isinstance(val_tot, pd.Series) else val_tot
                                        tot_str = str(val_tot).replace(",", "").strip()
                                        output["total_repayment"] = float(tot_str) if tot_str and tot_str != 'None' else None

                                    # Extract Interest Rate from Agreement Details table cells (e.g. '20 (17.82% (Y))' -> '17.82% per year')
                                    for c in df.columns:
                                        cell_val = row.get(c)
                                        cell_val = cell_val.iloc[0] if isinstance(cell_val, pd.Series) else cell_val
                                        cell_str = str(cell_val or '').strip()
                                        if '%' in cell_str:
                                            rate_m = re.search(r'(\d+(?:\.\d+)?)\s*%\s*(?:\(\s*Y\s*\)|per\s*year|\(y\))?', cell_str, re.I)
                                            if rate_m:
                                                output["interest_rate"] = f"{rate_m.group(1)}% per year"
                                                break
                                except (ValueError, TypeError, AttributeError) as e:
                                    logger.warning(f"Agreement table cell parse error: {e}")
                            break
                if output["loan_amount"] is not None:
                    break

        # 4b. Fallback Agreement Extraction via regex (used for OCR/scanned PDFs)
        if output["loan_amount"] is None:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                # ------------------------------------------------------------------
                # Fix 3: Only attempt OCR if the page is actually blank (scanned)
                #         and Tesseract is available — avoids silent skip
                # ------------------------------------------------------------------
                if not page_text.strip():
                    page_text = _ocr_page(page)
                if page_text:
                    match = re.search(
                        r'(\d{2}[-/]\d{2}[-/]\d{4})\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)\s+([\d,]+\.?\d*)',
                        page_text
                    )
                    if match:
                        output["loan_date"] = match.group(1)
                        try:
                            output["emi"] = float(match.group(2).replace(",", ""))
                        except Exception:
                            pass
                        output["loan_amount"] = float(match.group(4).replace(",", ""))
                        output["total_repayment"] = float(match.group(5).replace(",", ""))

                        # Total Installments (e.g. '4800000.00 6 (37.15% (Y)) 6')
                        after_tot = page_text[match.end():match.end() + 100]
                        inst_m = re.search(r'^\s*(\d+)', after_tot)
                        if inst_m:
                            try:
                                output["total_installments"] = int(inst_m.group(1))
                            except Exception:
                                pass

                        rate_m = re.search(r'(\d+(?:\.\d+)?)\s*%\s*(?:\(\s*Y\s*\)|per\s*year|\(y\))?', page_text, re.I)
                        if rate_m:
                            output["interest_rate"] = f"{rate_m.group(1)}% per year"
                        break

        # 4c. Secondary fallback for interest rate if not yet captured
        if not output.get("interest_rate"):
            for page in pdf.pages:
                p_text = page.extract_text() or ""
                rate_m = re.search(r'(\d+(?:\.\d+)?)\s*%\s*(?:\(\s*Y\s*\)|per\s*year|\(y\))', p_text, re.I)
                if rate_m:
                    output["interest_rate"] = f"{rate_m.group(1)}% per year"
                    break

        # 5. Amortization Schedule
        schedule = []
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            is_page_scanned = not page_text.strip()

            # ------------------------------------------------------------------
            # Fix 3: Consistent OCR via shared helper; only when page is scanned
            # ------------------------------------------------------------------
            if is_page_scanned:
                page_text = _ocr_page(page)

            if not is_page_scanned:
                tables = page.extract_tables()
                for table in tables:
                    df = pd.DataFrame(table)
                    header_index = None
                    for i, row in df.iterrows():
                        row_text = " ".join([str(x) for x in row if pd.notna(x)]).lower()
                        if "due" in row_text and "emi" in row_text and "interest" in row_text:
                            header_index = i
                            break

                    if header_index is not None:
                        df.columns = df.iloc[header_index]
                        df = df[header_index + 1:].reset_index(drop=True)
                        df.columns = [str(col).strip().replace("\n", " ") for col in df.columns]

                        if df.empty:
                            continue

                        for _, row in df.iterrows():
                            # ------------------------------------------------------------------
                            # Fix 1: Replace bare `except:` with `except Exception as e` + logging
                            # ------------------------------------------------------------------
                            try:
                                raw_date = str(row.get("Due Date") or row.get("Due") or "")
                                due_date = re.sub(r"\s+", "", raw_date)
                                emi = row.get("EMI")
                                emi = float(str(emi).replace(",", "").strip()) if emi and emi != 'None' else None
                                interest = row.get("Interest")
                                interest = float(str(interest).replace(",", "").strip()) if interest and interest != 'None' else None

                                if not due_date or emi is None:
                                    continue

                                schedule.append({
                                    "date": due_date,
                                    "amount": emi,
                                    "interest_amount": interest,
                                    "cheque_no": "",
                                    "received_date": None,
                                    "payment_date": None
                                })
                            except Exception as e:
                                logger.warning(f"Schedule row parse error (table): {e}")
                                continue
            else:
                # If it's scanned (OCR), extract_tables() won't work — fallback to OCR schedule parser
                if page_text.strip() and any(
                    kw in page_text for kw in ("Amortization", "Due Date", "EMI", "Due")
                ):
                    schedule_text = page_text
                    if "Amortization" in page_text:
                        schedule_text = page_text.split("Amortization", 1)[1]
                    elif "Due Date" in page_text:
                        schedule_text = page_text.split("Due Date", 1)[1]

                    parsed_ocr_rows = parse_ocr_schedule(schedule_text, expected_emi=output.get("emi"))
                    if parsed_ocr_rows:
                        schedule.extend(parsed_ocr_rows)
                    else:
                        matches = re.findall(
                            r'(?:\d+\s+)?(\d{2}[-/]\d{2}[-/]\d{4})\s+[\d,]+\.?\d*\s+([\d,]+\.?\d*)\s+[\d,]+\.?\d*',
                            schedule_text
                        )
                        for match in matches:
                            try:
                                schedule.append({
                                    "date": match[0],
                                    "amount": float(match[1].replace(",", "")),
                                    "interest_amount": float(match[2].replace(",", "")),
                                    "cheque_no": "",
                                    "received_date": None,
                                    "payment_date": None
                                })
                            except Exception as e:
                                logger.warning(f"Schedule row parse error (OCR regex): {e}")
                                continue

        output["table"] = schedule

    return output


def parse_docx_amount(val):
    """
    Parses amount expressions from docx documents.
    Supports:
      - 42.60L, 50L, 42.60 Lakh, 50 Lakhs
      - 1.5Cr, 1.5 Crore
      - 6,25,000.00, ₹6,25,000
    """
    if val is None:
        return None
    val_str = str(val).strip()
    if not val_str:
        return None
    # Check Crore
    cr_match = re.search(r'([\d,]+(?:\.\d+)?)\s*(?:cr|crore|crores)\b', val_str, re.I)
    if cr_match:
        num = float(cr_match.group(1).replace(',', ''))
        return round(num * 10000000.0, 2)
    # Check Lakh
    l_match = re.search(r'([\d,]+(?:\.\d+)?)\s*(?:l|lac|lacs|lakh|lakhs)\b', val_str, re.I)
    if l_match:
        num = float(l_match.group(1).replace(',', ''))
        return round(num * 100000.0, 2)
    # Standard numeric
    cleaned = re.sub(r'[^\d.]', '', val_str)
    try:
        return float(cleaned) if cleaned else None
    except ValueError:
        return None


def normalize_date_str(val):
    """
    Normalizes dates from text into DD-MM-YYYY format.
    """
    if not val:
        return ""
    m = re.search(r'(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})', str(val))
    if m:
        d, mth, y = m.groups()
        if len(y) == 2:
            y = "20" + y
        return f"{int(d):02d}-{int(mth):02d}-{y}"
    return str(val).strip()


def clean_entity_name(raw):
    """
    Cleans corporate/individual entity names:
    - Strips 'M/s.', 'M/S', 'M/s' prefix
    - Strips '(PAN No. ...)'
    - Strips trailing commas and whitespace
    """
    if not raw:
        return ""
    s = str(raw).strip()
    s = re.sub(r'^(?:M\s*/\s*s\.?|Messrs\.?)\s*', '', s, flags=re.I)
    s = re.sub(r'\s*\([Pp][Aa][Nn].*?\)', '', s)
    s = re.sub(r'[,.\s\-_/]+$', '', s)
    return s.strip()


def extract_docx_data(docx_path):
    """
    Parses loan details and repayment schedule from Word (.docx) documents.
    Supports both key-value pair format and letterhead/narrative agreement formats.
    Supports multi-page documents and multiple schedule tables.
    """
    if not DOCX_AVAILABLE:
        raise ValueError("python-docx is not installed on this system.")

    doc = docx.Document(docx_path)
    output = {
        "loan_number": None,
        "primary_account": None,
        "client_account": None,
        "address": None,
        "contact": None,
        "loan_date": None,
        "loan_amount": None,
        "total_repayment": None,
        "interest_rate": None,
        "emi": None,
        "total_installments": None,
        "table": []
    }

    full_text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    bank_info = []

    # 1. Parse simple KEY: VALUE paragraphs if present
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        parts = re.split(r'[:–—\-\u2013\u2014=]\s*', txt, maxsplit=1)
        if len(parts) == 2:
            k, v = parts[0].strip().upper(), parts[1].strip()
            if "REPAYMENT" in k and not output["total_repayment"]:
                output["total_repayment"] = parse_docx_amount(v)
            elif any(x in k for x in ["PAYMENT AMT", "PRINCIPAL", "LOAN AMT", "LOAN AMOUNT"]) and not output["loan_amount"]:
                output["loan_amount"] = parse_docx_amount(v)
            elif any(x in k for x in ["PARTY NAME", "CLIENT", "BORROWER", "CUSTOMER"]) and not output["client_account"]:
                output["client_account"] = clean_customer_name(clean_entity_name(v))
            elif any(x in k for x in ["PAYMENT DATE", "LOAN DATE", "DATE"]) and not output["loan_date"]:
                output["loan_date"] = normalize_date_str(v)
            elif any(x in k for x in ["ACCOUNT", "BRANCH"]) and not output["primary_account"]:
                output["primary_account"] = clean_entity_name(v).upper()
            elif any(x in k for x in ["LOAN NO", "LOAN NUMBER", "REF NO"]) and not output["loan_number"]:
                output["loan_number"] = v
            elif any(x in k for x in ["A/C", "IFSC", "BANK"]):
                bank_info.append(f"{parts[0].strip()}: {v}")
        elif any(x in txt.upper() for x in ["A/C", "IFSC", "BANK"]):
            bank_info.append(txt)

    # 2. Enhanced regex extraction for letter pad / narrative format
    # 2a. Loan Date
    if not output["loan_date"]:
        m_date = re.search(r'(?:Date|Loan\s+Date)\s*[:–—\-=]\s*(\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})', full_text, flags=re.I)
        if m_date:
            output["loan_date"] = normalize_date_str(m_date.group(1))

    # 2b. Client / Borrower Name
    if not output["client_account"]:
        m_client = re.search(r'(?:I\s*/\s*We|I/We)\s+(?:M\s*/\s*s\.?\s+)?([A-Z0-9\s&.,\'\u2019-]+?)\s+have\s+borrowed', full_text, flags=re.I)
        if not m_client:
            m_client = re.search(r'From\s+Name\s*:\s*(?:M\s*/\s*s\.?\s+)?([A-Z0-9\s&.,\'\u2019-]+?)(?:,\s*Bank|\s*Bank|\n|$)', full_text, flags=re.I)
        if m_client:
            output["client_account"] = clean_customer_name(clean_entity_name(m_client.group(1)))

    # 2c. Primary Account / Lender Name
    if not output["primary_account"]:
        m_primary = re.search(r'borrowed.*?from\s+(?:M\s*/\s*s\.?\s+)?([A-Z0-9\s&.,\'\u2019-]+?)(?:\.|\s*Towards|\n|$)', full_text, flags=re.I)
        if not m_primary:
            m_primary = re.search(r'To\s*(?:Date[^\n]*)?\n\s*(?:M\s*/\s*s\.?\s+)?([A-Z0-9\s&.,\'\u2019-]+?)(?:,\s*\(PAN|\s*\(PAN|,\s*Place|\n|$)', full_text, flags=re.I)
        if m_primary:
            output["primary_account"] = clean_entity_name(m_primary.group(1)).upper()

    # 2d. Loan Amount
    if not output["loan_amount"]:
        m_amt = re.search(r'loan\s+of\s+Rs\.?\s*([\d,]+(?:\.\d+)?)/?-?', full_text, flags=re.I)
        if not m_amt:
            m_amt = re.search(r'borrowed.*?Rs\.?\s*([\d,]+(?:\.\d+)?)/?-?', full_text, flags=re.I)
        if m_amt:
            output["loan_amount"] = parse_docx_amount(m_amt.group(1))

    # 2e. Bank / Contact info
    if not output["contact"]:
        m_bank = re.search(r'(?:From\s+Name:[^\n]*,\s*)?(Bank\s*:[^\n]+)', full_text, flags=re.I)
        if m_bank:
            output["contact"] = m_bank.group(1).strip()
        elif bank_info:
            output["contact"] = " | ".join(bank_info)

    # 2f. Address info
    if not output["address"]:
        m_branch = re.search(r'Branch\s*:\s*([^,\n]+(?:,\s*[^,\n]+)*?)(?:,\s*IFS|\s*IFS|\n|$)', full_text, flags=re.I)
        if m_branch:
            output["address"] = m_branch.group(1).strip()

    # 2g. Loan Number / Reference ID if mentioned
    if not output["loan_number"]:
        m_ref = re.search(r'(?:Loan\s*(?:No|Number|Id|ID|Ref)|Ref\s*(?:No|Number|Id)?)\s*[:–—\-=.]\s*([A-Za-z0-9\-_/]+)', full_text, flags=re.I)
        if m_ref:
            output["loan_number"] = m_ref.group(1).strip()

    # 3. Process tables across all pages
    total_from_table = None
    last_date_col = None
    last_amt_col = None
    last_chq_col = None

    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip().upper() for c in table.rows[0].cells]
        date_col = next((i for i, h in enumerate(header) if any(x in h for x in ["DATE", "DUE"])), None)
        amt_col = next((i for i, h in enumerate(header) if any(x in h for x in ["AMOUNT", "EMI", "INSTAL"])), None)
        chq_col = next((i for i, h in enumerate(header) if any(x in h for x in ["CHEQUE", "CHQ"])), None)

        if date_col is not None and amt_col is not None:
            last_date_col = date_col
            last_amt_col = amt_col
            last_chq_col = chq_col
            start_row = 1
        elif last_date_col is not None and last_amt_col is not None:
            # Multi-page continuation table without header
            date_col = last_date_col
            amt_col = last_amt_col
            chq_col = last_chq_col
            start_row = 0
        else:
            continue

        for row_idx in range(start_row, len(table.rows)):
            row = table.rows[row_idx]
            cells = [c.text.strip() for c in row.cells]
            # Check for header repeat
            if any("DATE" in c.upper() for c in cells) and any("AMOUNT" in c.upper() for c in cells):
                continue
            # Check for TOTAL / summary row
            if any("TOTAL" in c.upper() for c in cells):
                for c in cells:
                    parsed_c = parse_docx_amount(c)
                    if parsed_c and parsed_c > 0 and (total_from_table is None or parsed_c > total_from_table):
                        total_from_table = parsed_c
                continue

            row_date = normalize_date_str(cells[date_col]) if date_col is not None and date_col < len(cells) else ""
            row_amt = parse_docx_amount(cells[amt_col]) if amt_col is not None and amt_col < len(cells) else None
            row_chq = cells[chq_col] if chq_col is not None and chq_col < len(cells) else ""

            if row_date and row_amt is not None:
                output["table"].append({
                    "date": row_date,
                    "amount": row_amt,
                    "interest_amount": 0.0,
                    "cheque_no": row_chq,
                    "received_date": None,
                    "payment_date": None
                })

    # 4. Final calculations for installments, EMI, total repayment, and interest
    if output["table"]:
        output["total_installments"] = len(output["table"])
        output["emi"] = output["table"][0]["amount"]
        sum_amounts = sum(r["amount"] for r in output["table"])
        if not output["total_repayment"]:
            output["total_repayment"] = total_from_table if (total_from_table and total_from_table >= sum_amounts) else round(sum_amounts, 2)

        if (
            output["total_repayment"]
            and output["loan_amount"]
            and output["total_repayment"] >= output["loan_amount"]
        ):
            tot_int = round(output["total_repayment"] - output["loan_amount"], 2)
            n = len(output["table"])
            per_int = round(tot_int / n, 2)
            acc_int = 0.0
            for idx, r in enumerate(output["table"]):
                if idx == n - 1:
                    r["interest_amount"] = round(tot_int - acc_int, 2)
                else:
                    r["interest_amount"] = per_int
                    acc_int += per_int
            output["interest_rate"] = round((tot_int / output["loan_amount"]) * 100, 2)

    return output


def extract_document_data(file_path):
    """
    Dispatcher to extract loan data from either PDF or Word (.docx/.doc) document.
    """
    path_lower = str(file_path).lower()
    if path_lower.endswith(".docx") or path_lower.endswith(".doc"):
        return extract_docx_data(file_path)
    return _extract_pdf_data_original(file_path)


# Backward-compatible alias
extract_pdf_data = extract_document_data


def build_final_output(data, remaining_accounts):
    """
    Computes primary account amount, interest, shares, and secondary accounts breakdown
    based on extracted document data and configured secondary shares.
    """
    loan_amount = data.get("loan_amount")
    if loan_amount is None:
        raise ValueError("Loan amount not found in document")

    try:
        loan_amount = float(str(loan_amount).replace(",", "").strip())
    except (ValueError, TypeError):
        raise ValueError(f"Invalid loan amount: {loan_amount}")

    total_repayment = data.get("total_repayment")
    if total_repayment is not None:
        try:
            total_repayment = float(str(total_repayment).replace(",", "").strip())
        except (ValueError, TypeError):
            total_repayment = loan_amount
    else:
        total_repayment = loan_amount

    remaining_total = 0.0
    for acc in (remaining_accounts or []):
        try:
            share = float(str(acc.get("share", 0)).replace(",", "").strip())
        except (ValueError, TypeError):
            share = 0.0
        remaining_total += share

    if remaining_total > loan_amount:
        raise ValueError(f"Total secondary shares ({remaining_total}) exceed total loan amount ({loan_amount})")

    total_interest = round(total_repayment - loan_amount, 2) if total_repayment >= loan_amount else 0.0

    primary_amount = round(loan_amount - remaining_total, 2)
    primary_percentage = round((primary_amount / loan_amount) * 100, 2) if loan_amount > 0 else 0.0
    primary_interest = round((primary_amount / loan_amount) * total_interest, 2) if loan_amount > 0 else 0.0

    formatted_remaining = []
    for acc in (remaining_accounts or []):
        try:
            share = float(str(acc.get("share", 0)).replace(",", "").strip())
        except (ValueError, TypeError):
            share = 0.0
        pct = round((share / loan_amount) * 100, 2) if loan_amount > 0 else 0.0
        int_amt = round((share / loan_amount) * total_interest, 2) if loan_amount > 0 else 0.0
        tds = round(int_amt * 0.10, 2)
        formatted_remaining.append({
            "account_name": str(acc.get("name") or acc.get("account_name", "")).strip().upper(),
            "share": share,
            "percentage": pct,
            "interest_amount": int_amt,
            "tds": tds
        })

    return {
        "loan_number": data.get("loan_number") or "",
        "loan_date": data.get("loan_date") or "",
        "loan_amount": loan_amount,
        "total_repayment_amount": total_repayment,
        "total_interest": total_interest,
        "primary_account_name": str(data.get("primary_account") or "").strip().upper(),
        "primary_account_share": primary_percentage,
        "primary_account_amount": primary_amount,
        "primary_account_interest": primary_interest,
        "client_account_name": str(data.get("client_account") or "").strip(),
        "address": str(data.get("address") or "").strip(),
        "contact": str(data.get("contact") or "").strip(),
        "interest_rate": data.get("interest_rate") or "",
        "remaining_accounts": formatted_remaining,
        "table": data.get("table") or []
    }

